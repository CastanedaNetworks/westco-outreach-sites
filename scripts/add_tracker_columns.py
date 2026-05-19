"""Add Email and Instagram columns to existing outreach-tracker.xlsx in place.

Inserts after the Phone column, preserves all user-entered data, re-applies
header styling and column widths. Idempotent — safe to re-run.
"""
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "outreach-tracker.xlsx"

NEW_COLS = ["Email", "Instagram"]
INSERT_AFTER = "Phone"

wb = load_workbook(XLSX)
ws = wb.active

headers = [c.value for c in ws[1]]

# Idempotency — skip if already added
if all(c in headers for c in NEW_COLS):
    print("Columns already present; nothing to do.")
    raise SystemExit(0)

insert_idx = headers.index(INSERT_AFTER) + 2  # 1-based, after Phone
ws.insert_cols(insert_idx, amount=len(NEW_COLS))

# Header style — match existing
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor="2F5496")
thin = Side(border_style="thin", color="CCCCCC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
alt_fill = PatternFill("solid", fgColor="F2F2F2")

for offset, name in enumerate(NEW_COLS):
    col = insert_idx + offset
    c = ws.cell(row=1, column=col, value=name)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = border
    ws.column_dimensions[get_column_letter(col)].width = 26 if name == "Email" else 22

# Apply borders + alternating row shading + wrap to the new column cells
for row_idx in range(2, ws.max_row + 1):
    for offset in range(len(NEW_COLS)):
        col = insert_idx + offset
        cell = ws.cell(row=row_idx, column=col)
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if row_idx % 2 == 0:
            cell.fill = alt_fill

wb.save(XLSX)
print(f"Added columns {NEW_COLS} after '{INSERT_AFTER}'. Total cols: {ws.max_column}")

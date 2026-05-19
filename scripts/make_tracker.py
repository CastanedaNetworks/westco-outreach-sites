"""Build outreach tracking spreadsheet from prospects.csv + sites/ directory."""
import csv
import os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
SITES_DIR = ROOT / "sites"
CSV_PATH = ROOT / "data" / "prospects.csv"
OUT_PATH = ROOT / "outreach-tracker.xlsx"

# Load prospect data keyed by slug
prospects = {}
with open(CSV_PATH, encoding="utf-8") as f:
    for row in csv.DictReader(f):
        prospects[row["slug"]] = row

# Get the actual built sites
built_slugs = sorted(p.name for p in SITES_DIR.iterdir() if p.is_dir())

# Build rows
rows = []
for slug in built_slugs:
    p = prospects.get(slug, {})
    rows.append({
        "Business": p.get("name", slug.replace("-", " ").title()),
        "Category": p.get("category", ""),
        "City": p.get("city", ""),
        "Phone": p.get("phone", ""),
        "Email": "",
        "Instagram": "",
        "Rating": p.get("rating", ""),
        "Reviews": p.get("review_count", ""),
        "Site URL": p.get("generated_url", f"https://castanedanetworks.github.io/westco-outreach-sites/sites/{slug}/"),
        "Date Sent": "",
        "Sent Via": "",
        "Response?": "",
        "Response Date": "",
        "Follow-up Date": "",
        "Status": "",
        "Notes": "",
    })

# Build workbook
wb = Workbook()
ws = wb.active
ws.title = "Outreach"

headers = list(rows[0].keys())
ws.append(headers)

for r in rows:
    ws.append([r[h] for h in headers])

# Style header
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor="2F5496")
thin = Side(border_style="thin", color="CCCCCC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for col_idx, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=col_idx)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = border

# Column widths
widths = {
    "Business": 32, "Category": 18, "City": 14, "Phone": 16,
    "Email": 26, "Instagram": 22, "Rating": 8,
    "Reviews": 9, "Site URL": 70, "Date Sent": 12, "Sent Via": 12,
    "Response?": 11, "Response Date": 14, "Follow-up Date": 14,
    "Status": 18, "Notes": 40,
}
for i, h in enumerate(headers, 1):
    ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 14)

# Freeze header
ws.freeze_panes = "A2"

# Make URLs clickable
url_col = headers.index("Site URL") + 1
link_font = Font(color="0563C1", underline="single")
for row_idx in range(2, len(rows) + 2):
    cell = ws.cell(row=row_idx, column=url_col)
    if cell.value:
        cell.hyperlink = cell.value
        cell.font = link_font

# Data validation dropdowns
last_row = len(rows) + 1

def col_letter(name):
    return get_column_letter(headers.index(name) + 1)

dv_sent_via = DataValidation(type="list", formula1='"Email,Phone Call,Text,Instagram DM,In-Person,Mail"', allow_blank=True)
ws.add_data_validation(dv_sent_via)
dv_sent_via.add(f"{col_letter('Sent Via')}2:{col_letter('Sent Via')}{last_row}")

dv_response = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
ws.add_data_validation(dv_response)
dv_response.add(f"{col_letter('Response?')}2:{col_letter('Response?')}{last_row}")

dv_status = DataValidation(
    type="list",
    formula1='"Not Sent,Sent - Awaiting,Interested,Not Interested,No Response,Follow-up Needed,Converted,Dead"',
    allow_blank=True,
)
ws.add_data_validation(dv_status)
dv_status.add(f"{col_letter('Status')}2:{col_letter('Status')}{last_row}")

# Borders + alternating row shading for readability
alt_fill = PatternFill("solid", fgColor="F2F2F2")
for row_idx in range(2, last_row + 1):
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if row_idx % 2 == 0:
            if not cell.fill or cell.fill.fgColor.rgb in (None, "00000000"):
                cell.fill = alt_fill

ws.row_dimensions[1].height = 28

wb.save(OUT_PATH)
print(f"Wrote {OUT_PATH} with {len(rows)} rows")
